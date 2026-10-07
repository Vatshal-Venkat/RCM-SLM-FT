"""Text extraction for knowledge documents: Markdown, plain text, PDF and DOCX."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import yaml

SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".pdf", ".docx"}


@dataclass
class Section:
    heading: str | None
    text: str
    page: int | None = None


@dataclass
class Document:
    path: Path
    title: str
    sections: list[Section]
    metadata: dict = field(default_factory=dict)


_FRONTMATTER_RE = re.compile(r"^---\s*\n(.*?)\n---\s*\n", re.S)
_HEADING_RE = re.compile(r"^(#{1,4})\s+(.+?)\s*#*\s*$")


def load_document(path: Path) -> Document:
    suffix = path.suffix.lower()
    if suffix in (".md", ".markdown"):
        return _load_markdown(path)
    if suffix == ".txt":
        return Document(path, _title_from_name(path), [Section(None, path.read_text(encoding="utf-8"))])
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    raise ValueError(f"Unsupported document type: {path.suffix}")


def _title_from_name(path: Path) -> str:
    return path.stem.replace("_", " ").replace("-", " ").title()


def _load_markdown(path: Path) -> Document:
    raw = path.read_text(encoding="utf-8")
    meta: dict = {}
    m = _FRONTMATTER_RE.match(raw)
    if m:
        meta = yaml.safe_load(m.group(1)) or {}
        raw = raw[m.end():]

    title = meta.get("title")
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []  # heading path below the H1
    buf: list[str] = []

    def flush() -> None:
        text = "\n".join(buf).strip()
        if text:
            heading = " > ".join(h for _, h in stack) or None
            sections.append(Section(heading, text))
        buf.clear()

    for line in raw.splitlines():
        hm = _HEADING_RE.match(line)
        if hm:
            flush()
            level, heading = len(hm.group(1)), hm.group(2)
            if level == 1:
                title = title or heading
                stack.clear()
                continue
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, heading))
        else:
            buf.append(line)
    flush()
    return Document(path, title or _title_from_name(path), sections, meta)


def _load_pdf(path: Path) -> Document:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    title = (reader.metadata.title if reader.metadata and reader.metadata.title else None) or _title_from_name(path)
    sections = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(f"Page {i}", text, page=i))
    return Document(path, title, sections, {"pages": len(reader.pages)})


def _load_docx(path: Path) -> Document:
    import docx

    d = docx.Document(str(path))
    title = d.core_properties.title or _title_from_name(path)
    sections: list[Section] = []
    heading: str | None = None
    buf: list[str] = []
    for p in d.paragraphs:
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title":
            if buf:
                sections.append(Section(heading, "\n".join(buf)))
                buf = []
            heading = p.text.strip() or heading
        elif p.text.strip():
            buf.append(p.text.strip())
    if buf:
        sections.append(Section(heading, "\n".join(buf)))
    return Document(path, title, sections)
